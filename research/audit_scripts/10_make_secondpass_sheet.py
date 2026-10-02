"""Create the second labeling pass for the first-sentence unit (decisions D009 and D019).

The filter will classify the first line / first sentence of each output
(`rs_dic_llm.extraction.first_sentence`), while the first pass labeled the whole stored output.
Where the rule cuts words and the first-pass label was `definition` or `unclear`, the label of the
first sentence alone can differ (for example a question followed by a real definition), so those
rows are labeled again, on the first sentence as the model's entire answer. Rows the rule does
not cut keep their first-pass label; rows it cuts whose first-pass label was a non-definition
are taken to stay non-definitions (decision D019, option i).

Sheet: ID, Word, First sentence (as cut by the rule), Label, Notes. Rows shuffled, no model
names. Key (first-pass label, note, stored text, tab and row) goes to a separate file.
Refuses to overwrite existing files.
"""
import csv
import random
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
from rs_dic_llm.extraction import cut_removes_content, first_sentence  # noqa: E402

SOURCE = ROOT / "data" / "exp2_labeling_workbook_labeled.xlsx"
SHEET = ROOT / "data" / "exp2_secondpass_sheet.xlsx"
KEY = ROOT / "data" / "exp2_secondpass_KEY_do_not_open_until_done.csv"
SEED = "e2-secondpass-2026-10-01"
LABELS = "definition,echo,question,offtopic,restatement,refusal,unclear"
RELABEL = {"definition", "unclear"}

for path in (SHEET, KEY):
    if path.exists():
        sys.exit(f"{path.name} already exists; refusing to overwrite (it may hold your labels)")

source = openpyxl.load_workbook(SOURCE)
rows = []
for tab in [s for s in source.sheetnames if s != "Read me"]:
    ws = source[tab]
    for r in range(3, 103):
        stored = ws.cell(r, 2).value or ""
        label = ws.cell(r, 3).value
        if stored.strip() and cut_removes_content(stored) and label in RELABEL:
            rows.append(dict(tab=tab, row=r, word=ws.cell(r, 1).value, first_pass_label=label,
                             first_pass_note=ws.cell(r, 4).value or "", stored=stored.replace("\n", "\\n"),
                             shown=first_sentence(stored)))
random.Random(SEED).shuffle(rows)

wb = openpyxl.Workbook()
readme = wb.active
readme.title = "Read me"
lines = [
    ("Second pass: label the first sentence alone (E2, decisions D009 and D019)", True),
    ("", False),
    ("Each row shows only the FIRST SENTENCE of a model's output, as cut by the extraction rule. "
     "Label it as if it were the model's entire answer to the request to define the word: do not try "
     "to imagine what came after it, and do not look for the original output.", False),
    ("Use the same seven labels (dropdown) and the same rules as the first pass: "
     "research/E2_LABELING_INSTRUCTIONS.md sections 0-3 and 7. A question on its own is `question`; "
     "a lead-in that ends before the definition is `unclear` with the note `truncated`; a line that "
     "only repeats the task is `echo`.", False),
    (f"{len(rows)} rows, shuffled, model names hidden. They are the rows you labeled `definition` or "
     "`unclear` in the first pass whose output the rule shortens.", False),
    ("You can do this at any time; it does not need to wait a day. Do not open the key file "
     "(data/exp2_secondpass_KEY_do_not_open_until_done.csv) while you label.", False),
    ("When every row is labeled, save the file and tell the assistant; "
     "research/audit_scripts/11_secondpass_summary.py then compares the two passes.", False),
]
for i, (text, bold) in enumerate(lines, start=1):
    cell = readme.cell(i, 1, text)
    cell.font = Font(name="Arial", size=11, bold=bold)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
readme.column_dimensions["A"].width = 110

ws = wb.create_sheet("Label these")
fill = PatternFill("solid", fgColor="1F2937")
for col, (title, width) in enumerate([("ID", 6), ("Word", 18), ("First sentence", 80), ("Label", 14), ("Notes", 40)], start=1):
    cell = ws.cell(1, col, title)
    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell.fill = fill
    ws.column_dimensions[chr(64 + col)].width = width
for i, item in enumerate(rows, start=1):
    ws.cell(i + 1, 1, i)
    ws.cell(i + 1, 2, item["word"])
    ws.cell(i + 1, 3, item["shown"])
    for col in range(1, 6):
        ws.cell(i + 1, col).font = Font(name="Arial", size=11)
    ws.cell(i + 1, 3).alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
validation = DataValidation(type="list", formula1=f'"{LABELS}"', allow_blank=True, showErrorMessage=False)
validation.add(f"D2:D{len(rows) + 1}")
ws.add_data_validation(validation)
wb.save(SHEET)

with open(KEY, "w", newline="", encoding="utf-8-sig") as fh:
    writer = csv.DictWriter(fh, fieldnames=["id", "tab", "row", "word", "first_pass_label",
                                             "first_pass_note", "stored", "shown"])
    writer.writeheader()
    for i, item in enumerate(rows, start=1):
        writer.writerow(dict(id=i, **item))

by_tab = {}
for item in rows:
    by_tab[item["tab"]] = by_tab.get(item["tab"], 0) + 1
print(f"sheet: {SHEET.relative_to(ROOT)} ({len(rows)} rows)")
print(f"key  : {KEY.relative_to(ROOT)}")
print("rows per model:", by_tab)
print("first-pass labels:", {k: sum(1 for r in rows if r['first_pass_label'] == k) for k in RELABEL})
