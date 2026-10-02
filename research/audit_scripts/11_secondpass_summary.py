"""Compare the two labeling passes and build the validation set for the first-sentence unit.

Usage (defaults are the real files):
    python research/audit_scripts/11_secondpass_summary.py
    python research/audit_scripts/11_secondpass_summary.py --sheet SHEET.xlsx --key KEY.csv

1. How many of the rows labeled `definition` in the first pass are still `definition` when only
   the first sentence is judged (the rows the strict extraction rule can lose), by note tag and
   by model, with the flipped rows listed.
2. The validation set for the first-sentence unit (all 2,300 labeled rows):
   - rows the rule does not cut keep their first-pass label (source first_pass);
   - rows it cuts whose first-pass label was `definition` or `unclear` take the second-pass
     label (source second_pass);
   - rows it cuts whose first-pass label was a non-definition are taken to stay
     non-definitions, class not re-labeled (source inferred; decision D019, option i).
   Written to outputs/11_unit_a_validation_set.csv.
"""
import argparse
import collections
import csv
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
from rs_dic_llm.extraction import cut_removes_content, first_sentence  # noqa: E402

LABELS = ["definition", "echo", "question", "offtopic", "restatement", "refusal", "unclear"]
OUT = Path(__file__).resolve().parent / "outputs" / "11_unit_a_validation_set.csv"
sys.stdout.reconfigure(encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--sheet", default=str(ROOT / "data" / "exp2_secondpass_sheet.xlsx"))
    parser.add_argument("--key", default=str(ROOT / "data" / "exp2_secondpass_KEY_do_not_open_until_done.csv"))
    parser.add_argument("--workbook", default=str(ROOT / "data" / "exp2_labeling_workbook_labeled.xlsx"))
    args = parser.parse_args()

    key = {int(r["id"]): r for r in csv.DictReader(open(args.key, encoding="utf-8-sig"))}
    ws = openpyxl.load_workbook(args.sheet)["Label these"]
    second = {}
    for r in range(2, ws.max_row + 1):
        rid, label = ws.cell(r, 1).value, ws.cell(r, 4).value
        if rid is not None:
            second[int(rid)] = (str(label).strip().lower() if label not in (None, "") else None, ws.cell(r, 5).value or "")
    missing = [i for i in key if second.get(i, (None,))[0] is None]
    invalid = [(i, second[i][0]) for i in key if second.get(i, (None,))[0] not in (None, *LABELS)]
    print(f"second-pass rows {len(key)}; unlabeled {len(missing)} {missing[:10]}; invalid {invalid}")
    if missing or invalid:
        sys.exit("fix the sheet first")

    by_loc = {(k["tab"], int(k["row"])): (second[int(i)][0], second[int(i)][1], k) for i, k in key.items()}
    defs = [(i, k) for i, k in key.items() if k["first_pass_label"] == "definition"]
    kept = [i for i, k in defs if second[i][0] == "definition"]
    print(f"\n1. ROWS LABELED `definition` IN THE FIRST PASS AND CUT BY THE RULE (n={len(defs)})")
    print(f"   still `definition` on the first sentence alone: {len(kept)} ({100 * len(kept) / len(defs):.1f}%); "
          f"changed: {len(defs) - len(kept)}")
    print("   where they went:", dict(collections.Counter(second[i][0] for i, _ in defs if second[i][0] != "definition")))
    print("   by first-pass note tag (rows / still definition):")
    tags = collections.defaultdict(lambda: [0, 0])
    for i, k in defs:
        tag = (k["first_pass_note"].split(";")[0].strip() or "(no note)")
        tags[tag][0] += 1
        tags[tag][1] += second[i][0] == "definition"
    for tag, (n, s) in sorted(tags.items(), key=lambda x: -x[1][0]):
        print(f"     {tag:26} {n:3} / {s:3}")
    print("   by model (rows / still definition):")
    models = collections.defaultdict(lambda: [0, 0])
    for i, k in defs:
        models[k["tab"]][0] += 1
        models[k["tab"]][1] += second[i][0] == "definition"
    for tab, (n, s) in models.items():
        print(f"     {tab:16} {n:3} / {s:3}")
    print("   changed rows:")
    for i, k in defs:
        if second[i][0] != "definition":
            print(f"     id {i:2} [{k['tab']} row {k['row']}, {k['word']}] -> {second[i][0]}: {k['shown'][:90]!r}")
    uncl = [(i, k) for i, k in key.items() if k["first_pass_label"] == "unclear"]
    print("   first-pass `unclear` rows now:", {k["tab"] + " " + k["word"]: second[i][0] for i, k in uncl})

    print("\n2. VALIDATION SET FOR THE FIRST-SENTENCE UNIT")
    wb = openpyxl.load_workbook(args.workbook)
    rows = []
    for tab in [s for s in wb.sheetnames if s != "Read me"]:
        w = wb[tab]
        for r in range(3, 103):
            stored = w.cell(r, 2).value or ""
            first = w.cell(r, 3).value
            cut = bool(stored.strip()) and cut_removes_content(stored)
            if not cut:
                label, source = first, "first_pass"
            elif (tab, r) in by_loc:
                label, source = by_loc[(tab, r)][0], "second_pass"
            else:
                label, source = "non-definition (class not re-labeled)", "inferred"
            rows.append(dict(tab=tab, row=r, word=w.cell(r, 1).value, unit_a_text=first_sentence(stored),
                             unit_a_label=label, source=source, first_pass_label=first))
    OUT.parent.mkdir(exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    print(f"   {len(rows)} rows by source: {dict(collections.Counter(x['source'] for x in rows))}")
    for group, test in (("base models", lambda t: t.endswith("-pt")), ("all others", lambda t: not t.endswith("-pt"))):
        sub = [x for x in rows if test(x["tab"])]
        first = sum(1 for x in sub if x["first_pass_label"] == "definition") / len(sub)
        unit = sum(1 for x in sub if x["unit_a_label"] == "definition") / len(sub)
        print(f"   share `definition`, {group}: whole output {100 * first:.1f}% -> first sentence {100 * unit:.1f}%")
    print(f"   wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
