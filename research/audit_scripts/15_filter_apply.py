"""Apply the non-definition filter (D023) to every stored record and count what it drops.

Usage:
    python research/audit_scripts/15_filter_apply.py                       # drop rates, every set and model
    python research/audit_scripts/15_filter_apply.py --audit MODEL --set main --kind drops --n 40
    python research/audit_scripts/15_filter_apply.py --audit MODEL --set e5 --kind keeps --n 40

Records are read as stored (`definition`), cut to the first sentence with rs_dic_llm.extraction
(the primary unit) and classified with rs_dic_llm.quality_filter. Nothing is written back to the
definitions. `--audit` prints random drops or keeps for one model, leaving out the 100 hand-labeled
words so that the test words are not looked at: it is a read-through for bugs on records that have
no label, not a validation.

The table is written to outputs/15_filter_drop_rates.csv. The mapping from set name to files is the
one used by 13_controls_vs_main.py.
"""
import argparse
import collections
import csv
import json
import random
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
sys.stdout.reconfigure(encoding="utf-8")
from rs_dic_llm.extraction import first_sentence  # noqa: E402
from rs_dic_llm.quality_filter import classify  # noqa: E402

OUT = Path(__file__).resolve().parent / "outputs"
WORKBOOK = ROOT / "data" / "exp2_labeling_workbook_labeled.xlsx"
PATHS = {
    "main": "data/definitions/{}_42.jsonl",
    "e3": "results_2026-09-09_Runpod/results_e3 (fewshot)/definitions/{}_42.jsonl",
    "e4": "results_2026-09-09_Runpod/results_e4 (lenght)/definitions/{}_42.jsonl",
    "e5": "results_2026-09-09_Runpod/results_e5 (no template)/definitions/{}_42.jsonl",
}
REASONS = ("fragment", "question", "echo", "refusal", "offtopic")


def models_and_labeled_words():
    wb = openpyxl.load_workbook(WORKBOOK)
    tabs = [s for s in wb.sheetnames if s != "Read me"]
    words = {str(wb[tabs[0]].cell(r, 1).value).strip().lower() for r in range(3, 103) if wb[tabs[0]].cell(r, 1).value}
    return tabs, words


def records(set_name, model):
    path = ROOT / PATHS[set_name].format(model)
    if not path.exists():
        return None
    return [json.loads(line) for line in open(path, encoding="utf-8") if line.strip()]


def decide(record, **kw):
    stored = record.get("definition") or ""
    text = first_sentence(stored)
    return stored, text, classify(text, record["word"], **kw)


def table(models, **kw):
    rows = []
    for set_name in PATHS:
        for model in models:
            recs = records(set_name, model)
            if recs is None:
                continue
            counts = collections.Counter()
            rules = collections.Counter()
            for r in recs:
                _, _, d = decide(r, **kw)
                if not d.keep:
                    counts[d.reason] += 1
                    rules[d.rule] += 1
            n = len(recs)
            dropped = sum(counts.values())
            rows.append(dict(set=set_name, model=model, n=n, dropped=dropped, dropped_pct=round(100 * dropped / n, 1),
                             **{r: counts[r] for r in REASONS}, top_rule=rules.most_common(1)[0][0] if rules else ""))
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--audit")
    parser.add_argument("--set", default="main", choices=list(PATHS))
    parser.add_argument("--kind", default="drops", choices=("drops", "keeps"))
    parser.add_argument("--n", type=int, default=40)
    parser.add_argument("--drop-imperatives", action="store_true")
    args = parser.parse_args()
    models, labeled = models_and_labeled_words()
    kw = dict(drop_imperatives=args.drop_imperatives)

    if args.audit:
        recs = records(args.set, args.audit)
        if recs is None:
            sys.exit(f"no file for {args.set} / {args.audit}")
        pool = []
        for r in recs:
            if str(r["word"]).strip().lower() in labeled:
                continue
            stored, text, d = decide(r, **kw)
            if (not d.keep) == (args.kind == "drops"):
                pool.append((r["word"], d.reason, d.rule, text))
        random.Random("audit-" + args.set + args.audit).shuffle(pool)
        print(f"{args.set} / {args.audit}: {len(pool)} unlabeled-word records with the filter decision `{args.kind}`; showing {min(args.n, len(pool))}")
        for word, reason, rule, text in pool[: args.n]:
            print(f"   [{str(word)[:12]:12}] {reason}/{rule}: {text[:120]!r}")
        return

    rows = table(models, **kw)
    OUT.mkdir(exist_ok=True)
    name = "15_filter_drop_rates_imperatives.csv" if args.drop_imperatives else "15_filter_drop_rates.csv"
    with open(OUT / name, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    print(f"{'set':5}{'model':26}{'n':>6}{'dropped':>9}{'%':>7}  " + "".join(f"{r:>10}" for r in REASONS) + "  top rule")
    for r in rows:
        print(f"{r['set']:5}{r['model']:26}{r['n']:>6}{r['dropped']:>9}{r['dropped_pct']:>7.1f}  "
              + "".join(f"{r[x]:>10}" for x in REASONS) + f"  {r['top_rule']}")
    print(f"\nwrote {(OUT / name).relative_to(ROOT)}")


if __name__ == "__main__":
    main()
