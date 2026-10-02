"""Create the E3 check sheet and its key (decision D027).

The E3 result (few-shot base models reach the R of the instruct models) rests on E3 outputs that the filter
keeps almost entirely and that were never hand-labeled. The researcher labels 100 of them, blind to model and to
the filter, with the seven labels of the first pass; 22_e3_check_summary.py then reads the labels by the rule fixed in
D027.

Sample: 25 words drawn at random (seed fixed below) from the words that are not among the 100 hand-labeled words and
whose E3 record is usable in all four models; the same 25 words for Gemma3-1B-pt, 4B-pt, 12B-pt and 27B-pt, so 100
rows, shuffled. The text shown is the first sentence of the stored definition (rs_dic_llm.extraction), the text that
enters the graph. Sheet columns: ID, Word, Model Output, Label (dropdown of the seven labels), Notes. The key (model,
stored text, and the decision of the default filter for each row) goes to a separate file that must not be opened
before labeling is finished.

Refuses to overwrite existing files, so a filled sheet can never be lost by re-running.
"""
import csv
import json
import random
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
sys.stdout.reconfigure(encoding="utf-8")
from rs_dic_llm.extraction import first_sentence  # noqa: E402
from rs_dic_llm.quality_filter import classify  # noqa: E402

SHEET = ROOT / "data" / "exp2_e3_check_sheet.xlsx"
KEY = ROOT / "data" / "exp2_e3_check_KEY_do_not_open_until_done.csv"
WORKBOOK = ROOT / "data" / "exp2_labeling_workbook_labeled.xlsx"
E3 = ROOT / "results_2026-09-09_Runpod" / "results_e3 (fewshot)" / "definitions"
MODELS = ["Gemma3-1B-pt", "Gemma3-4B-pt", "Gemma3-12B-pt", "Gemma3-27B-pt"]
SEED = "e3-check-2026-10-02"
N_WORDS = 25
LABELS = "definition,echo,question,offtopic,restatement,refusal,unclear"
USABLE = {"ok", "self_referential"}

for path in (SHEET, KEY):
    if path.exists():
        sys.exit(f"{path.name} already exists; refusing to overwrite (it may hold your labels)")

wb_src = openpyxl.load_workbook(WORKBOOK, read_only=True)
tab = [s for s in wb_src.sheetnames if s != "Read me"][0]
labeled = {str(r[0]).strip().lower() for r in wb_src[tab].iter_rows(min_row=3, max_row=102, max_col=1, values_only=True) if r[0]}

records = {}
for model in MODELS:
    for line in open(E3 / f"{model}_42.jsonl", encoding="utf-8"):
        if line.strip():
            r = json.loads(line)
            records[(model, str(r["word"]).strip().lower())] = r
words = sorted({w for (_, w) in records})
candidates = [w for w in words if w not in labeled and all(records.get((m, w), {}).get("status") in USABLE for m in MODELS)]
rng = random.Random(SEED)
chosen = sorted(rng.sample(candidates, N_WORDS))
picked = [(m, w) for m in MODELS for w in chosen]
rng.shuffle(picked)

wb = openpyxl.Workbook()
readme = wb.active
readme.title = "Read me"
lines = [
    ("E3 check (E2): are the few-shot base outputs definitions?", True),
    ("", False),
    ("Purpose: the E3 result (few-shot base models reach the reciprocity of the instruction-tuned models) rests on these "
     "outputs being real definitions, and the filter has never been checked on few-shot outputs. You label 100 of them, blind "
     "to the model and to the filter.", False),
    ("Use the same seven labels (dropdown) and the same rules as the first pass: research/E2_LABELING_INSTRUCTIONS.md sections "
     "0-3 and 7. Judge the whole text as shown. Notes are optional; use the same tags if you add any.", False),
    ("The rule that decides what the result means was fixed before you label: research/DECISION_LOG.md, D027.", False),
    ("Do not open the key file (data/exp2_e3_check_KEY_do_not_open_until_done.csv). Do not reorder or delete rows. "
     "The same 25 words appear four times, with different outputs; the model is not shown on purpose.", False),
    ("Time: about 30 minutes. There is no waiting period: these outputs were never labeled before.", False),
    ("When you have labeled every row, save the file and tell the assistant.", False),
]
for i, (text, bold) in enumerate(lines, start=1):
    cell = readme.cell(i, 1, text)
    cell.font = Font(name="Arial", size=11, bold=bold)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
readme.column_dimensions["A"].width = 110

ws = wb.create_sheet("Label these")
fill = PatternFill("solid", fgColor="1F2937")
for col, (title, width) in enumerate([("ID", 6), ("Word", 18), ("Model Output", 80), ("Label", 14), ("Notes", 40)], start=1):
    cell = ws.cell(1, col, title)
    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell.fill = fill
    ws.column_dimensions[chr(64 + col)].width = width
key_rows = []
for i, (model, word) in enumerate(picked, start=1):
    rec = records[(model, word)]
    stored = rec.get("definition") or ""
    shown = first_sentence(stored)
    decision = classify(shown, rec["word"])
    r = i + 1
    ws.cell(r, 1, i)
    ws.cell(r, 2, rec["word"])
    ws.cell(r, 3, shown)
    for col in (1, 2, 3, 4, 5):
        ws.cell(r, col).font = Font(name="Arial", size=11)
    ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
    key_rows.append(dict(id=i, model=model, word=rec["word"], stored=stored, shown=shown,
                         filter_keep=decision.keep, filter_reason=decision.reason, filter_rule=decision.rule))
ws.freeze_panes = "A2"
validation = DataValidation(type="list", formula1=f'"{LABELS}"', allow_blank=True, showErrorMessage=False)
validation.add(f"D2:D{len(picked) + 1}")
ws.add_data_validation(validation)
wb.save(SHEET)

with open(KEY, "w", newline="", encoding="utf-8-sig") as fh:
    writer = csv.DictWriter(fh, fieldnames=list(key_rows[0].keys()))
    writer.writeheader()
    writer.writerows(key_rows)

print(f"sheet: {SHEET.relative_to(ROOT)} ({len(picked)} rows: {N_WORDS} words x {len(MODELS)} models, {len(candidates)} candidate words)")
print(f"key  : {KEY.relative_to(ROOT)} (do not open before labeling is finished)")
print(f"rows the default filter drops: {sum(1 for k in key_rows if not k['filter_keep'])} of {len(key_rows)} (not shown to the labeler)")
