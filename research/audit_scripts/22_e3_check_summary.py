"""Read the E3 hand check by the rule of decision D027 (run it once, after the sheet is fully labeled).

Reads data/exp2_e3_check_sheet.xlsx (labels) and data/exp2_e3_check_KEY_do_not_open_until_done.csv (model and the default
filter's decision per row). Prints: the share of `definition` among the outputs the default filter keeps, overall and per
model, with the exact 95 % interval; the rule of D027 applied ("E3 stands" if at least 90 % of the kept outputs are
`definition` and no model is below 80 % of its kept outputs); the agreement of the filter with the labels on this
sample; the kept outputs that are not definitions and the dropped ones that are.
Run once, from the repository root, after the labeling:
    python research/audit_scripts/22_e3_check_summary.py | tee research/audit_scripts/outputs/22_e3_check_summary.txt
(two optional arguments, a sheet path and a key path, exist only for testing on synthetic files).
"""
import collections
import csv
import sys
from pathlib import Path

import openpyxl
from scipy.stats import beta

ROOT = Path(__file__).resolve().parents[2]
sys.stdout.reconfigure(encoding="utf-8")
SHEET = Path(sys.argv[1]) if len(sys.argv) > 2 else ROOT / "data" / "exp2_e3_check_sheet.xlsx"  # paths only for testing on synthetic files
KEY = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "data" / "exp2_e3_check_KEY_do_not_open_until_done.csv"
LABELS = {"definition", "echo", "question", "offtopic", "restatement", "refusal", "unclear"}
MODELS = ["Gemma3-1B-pt", "Gemma3-4B-pt", "Gemma3-12B-pt", "Gemma3-27B-pt"]
OVERALL_MIN, MODEL_MIN = 0.90, 0.80


def cp(k, n, level=0.95):
    a = (1 - level) / 2
    return (0.0 if k == 0 else beta.ppf(a, k, n - k + 1)), (1.0 if k == n else beta.ppf(1 - a, k + 1, n - k))


key = {int(r["id"]): r for r in csv.DictReader(open(KEY, encoding="utf-8-sig"))}
ws = openpyxl.load_workbook(SHEET)["Label these"]
label = {}
for r in range(2, ws.max_row + 1):
    rid, lab = ws.cell(r, 1).value, ws.cell(r, 4).value
    if rid is not None:
        label[int(rid)] = str(lab).strip().lower() if lab not in (None, "") else None
missing = [i for i in key if label.get(i) is None]
invalid = [(i, label[i]) for i in key if label.get(i) not in (None, *LABELS)]
print(f"rows {len(key)}; unlabeled {len(missing)} {missing[:10]}; invalid {invalid}")
if missing or invalid:
    sys.exit("fix the sheet first")

rows = [dict(id=i, model=key[i]["model"], word=key[i]["word"], text=key[i]["shown"], label=label[i], keep=key[i]["filter_keep"] == "True") for i in sorted(key)]
print("\n1. LABELS OF ALL 100 OUTPUTS, per model")
for m in MODELS:
    c = collections.Counter(r["label"] for r in rows if r["model"] == m)
    print(f"   {m:16} " + ", ".join(f"{k} {c[k]}" for k in sorted(c, key=lambda k: -c[k])))

kept = [r for r in rows if r["keep"]]
print(f"\n2. SHARE OF `definition` AMONG THE OUTPUTS THE DEFAULT FILTER KEEPS ({len(kept)} of {len(rows)} kept)")
d_all = sum(r["label"] == "definition" for r in kept)
lo, hi = cp(d_all, len(kept))
print(f"   overall: {d_all} of {len(kept)} = {100 * d_all / len(kept):.1f} %  (95 % interval {100 * lo:.0f}-{100 * hi:.0f} %)")
per_model_ok = True
for m in MODELS:
    km = [r for r in kept if r["model"] == m]
    d = sum(r["label"] == "definition" for r in km)
    lo, hi = cp(d, len(km)) if km else (0, 0)
    share = d / len(km) if km else 0.0
    per_model_ok &= share >= MODEL_MIN
    print(f"   {m:16} {d} of {len(km)} = {100 * share:.0f} %  (95 % interval {100 * lo:.0f}-{100 * hi:.0f} %)")
stands = d_all / len(kept) >= OVERALL_MIN and per_model_ok
print(f"\n3. RULE OF D027: at least {100 * OVERALL_MIN:.0f} % of the kept outputs are definitions and no model below {100 * MODEL_MIN:.0f} %")
print("   ->", "E3 STANDS" if stands else "RULE NOT MET: the E3 section is worded as unvalidated and D026 is reopened with the advisor")

tp = sum(1 for r in rows if not r["keep"] and r["label"] != "definition")
fp = sum(1 for r in rows if not r["keep"] and r["label"] == "definition")
fn = sum(1 for r in rows if r["keep"] and r["label"] != "definition")
tn = sum(1 for r in rows if r["keep"] and r["label"] == "definition")
print(f"\n4. THE FILTER ON THIS SAMPLE: dropped and not a definition {tp}, dropped but a definition {fp}, kept but not a definition {fn}, kept and a definition {tn}")
print("\n5. KEPT OUTPUTS THAT ARE NOT DEFINITIONS")
for r in rows:
    if r["keep"] and r["label"] != "definition":
        print(f"   [{r['model']:14} {r['word'][:12]:12}] {r['label']:11} {r['text'][:120]!r}")
print("\n6. DROPPED OUTPUTS THAT ARE DEFINITIONS")
for r in rows:
    if not r["keep"] and r["label"] == "definition":
        print(f"   [{r['model']:14} {r['word'][:12]:12}] {r['text'][:120]!r}")
