"""Review of the E2 first-pass hand labels (data/exp2_labeling_workbook_labeled.xlsx).

Audit script (2026-10-01). Checks completeness and integrity, summarises labels and notes,
and lists rows whose label conflicts with simple surface features or with how another model's
identical output was labeled, so the researcher can re-check them.

This is a label-quality review, not the non-definition filter: nothing here is used to
classify outputs, and the surface features are deliberately crude. Flagged rows are
suggestions to look at, not corrections.

Writes outputs/5_e2_label_flags.csv (one row per flag).
"""
import collections
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "data" / "exp2_labeling_workbook_labeled.xlsx"
OUT_CSV = Path(__file__).resolve().parent / "outputs" / "5_e2_label_flags.csv"
VALID = ["definition", "echo", "question", "offtopic", "restatement", "refusal", "unclear"]
PROMPT_PHRASES = ["short sentence", "common english words", "do not use the word",
                  "define the noun", "define the verb", "define the adjective", "in one short"]

sys.stdout.reconfigure(encoding="utf-8")
wb = openpyxl.load_workbook(WORKBOOK)
tabs = [n for n in wb.sheetnames if n != "Read me"]


def records(tab: str) -> dict:
    by = collections.defaultdict(list)
    path = ROOT / "data" / "definitions" / f"{tab}_42.jsonl"
    for line in open(path, encoding="utf-8"):
        if line.strip():
            r = json.loads(line)
            by[r["lemma"]].append(r)
    return by


rows = []  # one dict per labeled row
for tab in tabs:
    ws = wb[tab]
    src = records(tab)
    for r in range(3, 103):
        word, out, label, note = (ws.cell(r, c).value for c in (1, 2, 3, 4))
        rows.append(dict(tab=tab, row=r, word=word, output=out or "", label=label,
                         note=note or "",
                         in_source=any((x["definition"] or "") == (out or "") for x in src[word])))

is_base = lambda t: t.endswith("-pt")

# ── A. completeness and integrity ────────────────────────────────────────────
print("A. COMPLETENESS AND INTEGRITY")
print(f"  model tabs {len(tabs)}; rows {len(rows)}; labeled {sum(1 for x in rows if x['label'])};"
      f" invalid {[ (x['tab'], x['row'], x['label']) for x in rows if x['label'] not in VALID]}")
print(f"  rows whose output is not a stored record of that word: {sum(1 for x in rows if not x['in_source'])}")
words_per_tab = {t: [x["word"] for x in rows if x["tab"] == t] for t in tabs}
print(f"  same 100 words in the same order in every tab: {len({tuple(v) for v in words_per_tab.values()}) == 1}")

# ── B. label distribution ────────────────────────────────────────────────────
print("\nB. LABELS PER MODEL (counts of 100)")
hdr = f"  {'model':24}" + "".join(f"{l[:6]:>8}" for l in VALID)
print(hdr)
dist = {t: collections.Counter(x["label"] for x in rows if x["tab"] == t) for t in tabs}
for t in tabs:
    print(f"  {t + (' (base)' if is_base(t) else ''):24}" + "".join(f"{dist[t][l]:>8}" for l in VALID))
groups = {
    "Gemma3 instruct (5)": [t for t in tabs if t.startswith("Gemma3") and not is_base(t)],
    "Gemma3 instruct, no 270M": [t for t in tabs if t.startswith("Gemma3") and not is_base(t) and t != "Gemma3-270M"],
    "Gemma3 base (5)": [t for t in tabs if is_base(t)],
    "Qwen2.5 (7)": [t for t in tabs if t.startswith("Qwen2.5")],
    "Qwen3 (6)": [t for t in tabs if t.startswith("Qwen3")],
}
print("\n  group totals (percent of rows)")
for g, ts in groups.items():
    tot = sum(sum(dist[t].values()) for t in ts)
    print(f"  {g:26}" + "".join(f"{l[:6]:>8}" for l in VALID) + f"   n={tot}")
    print(f"  {'':26}" + "".join(f"{100 * sum(dist[t][l] for t in ts) / tot:>7.1f}%" for l in VALID))

# ── C. notes ─────────────────────────────────────────────────────────────────
print("\nC. NOTES")
noted = [x for x in rows if x["note"]]
print(f"  rows with a note: {len(noted)}")
tag = collections.Counter()
for x in noted:
    for t in re.split(r"\s*(?:\+|;)\s*", x["note"].strip().lower()):
        if t:
            tag[t] += 1
print("  tags (split on + and ;), 6 or more uses:",
      ", ".join(f"{t} {c}" for t, c in tag.most_common() if c >= 6))

# ── D. consistency flags ─────────────────────────────────────────────────────
flags = []


def flag(x, reason):
    flags.append(dict(reason=reason, tab=x["tab"], row=x["row"], word=x["word"],
                      label=x["label"], note=x["note"], output=x["output"].replace("\n", "\\n")[:200]))


norm = lambda s: re.sub(r"\s+", " ", s.strip().lower())
toks = lambda s: re.findall(r"[a-z'-]+", s.lower())

# D1 identical output for the same word in several models but different labels
groups_ = collections.defaultdict(list)
for x in rows:
    if x["output"].strip():
        groups_[(x["word"], norm(x["output"]))].append(x)
d1 = 0
for g in groups_.values():
    if len({x["tab"] for x in g}) >= 2 and len({x["label"] for x in g}) >= 2:
        d1 += 1
        for x in g:
            flag(x, "D1 identical output to another model but different label")
for x in rows:
    out, lab, w = x["output"], x["label"], (x["word"] or "").lower()
    n_words = len(out.split())
    if lab == "definition" and "?" in out and "q+def" not in x["note"].lower() and "q" not in x["note"].lower():
        flag(x, "D2 labeled definition but output contains '?' and no q+def note")
    if lab == "question" and "?" not in out:
        flag(x, "D3 labeled question but output has no '?'")
    if lab == "echo" and not any(p in out.lower() for p in PROMPT_PHRASES):
        flag(x, "D4 labeled echo but no prompt phrase in output")
    if lab == "offtopic" and not is_base(x["tab"]) and w in toks(out) and n_words <= 25:
        flag(x, "D5 instruct model, offtopic, but output is short and contains the target word")
    if lab == "restatement" and n_words > 4:
        flag(x, "D6 labeled restatement but output longer than 4 words")
    if lab == "definition" and n_words <= 2:
        flag(x, "D7 labeled definition but output has 2 words or fewer")
    if not out.strip() and lab not in ("unclear",):
        flag(x, "D8 blank output not labeled unclear")
    if lab == "unclear":
        flag(x, "D9 unclear (list for a second look)")
print(f"\nD. CONSISTENCY FLAGS (crude surface checks; {len(flags)} flags)")
by_reason = collections.Counter(f["reason"] for f in flags)
for k, v in sorted(by_reason.items()):
    print(f"  {v:4}  {k}")
print(f"  (D1 groups with conflicting labels: {d1}; total identical-output groups over several models: "
      f"{sum(1 for g in groups_.values() if len({x['tab'] for x in g}) >= 2)})")

OUT_CSV.parent.mkdir(exist_ok=True)
with open(OUT_CSV, "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.DictWriter(fh, fieldnames=["reason", "tab", "row", "word", "label", "note", "output"])
    w.writeheader()
    w.writerows(sorted(flags, key=lambda f: (f["reason"], f["tab"], f["row"])))
print(f"\nwrote {OUT_CSV.relative_to(ROOT)}")
