"""What the uniform first-line / first-sentence extraction does to the stored definitions.

Audit script (2026-10-01) for decision D009. Applies rs_dic_llm.extraction.first_sentence to
every stored definition of the 23 paper models and of E3, E4 and E5, and reports how often the
text changes, how often words are removed (not just markup or list markers), which part of the
rule does the cutting, and what happens to the graphs. It also lists the labeled rows whose
content the rule would cut: these are the rows of the second labeling pass.

Writes outputs/7_second_pass_rows.csv.
"""
import collections
import csv
import json
import os
import random
import sys
from multiprocessing import Pool
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
os.chdir(ROOT)
sys.stdout.reconfigure(encoding="utf-8")

from rs_dic_llm.config import MODEL_REGISTRY  # noqa: E402
from rs_dic_llm.extraction import (  # noqa: E402
    _cleaned_lines, _cut_sentence, cut_removes_content, first_sentence)
from rs_dic_llm.graph_build import build_graph  # noqa: E402

NAMES = [display for display, *_ in MODEL_REGISTRY]
PATHS = {
    "main": "data/definitions/{}_42.jsonl",
    "E3": "results_2026-09-09_Runpod/results_e3 (fewshot)/definitions/{}_42.jsonl",
    "E4": "results_2026-09-09_Runpod/results_e4 (lenght)/definitions/{}_42.jsonl",
    "E5": "results_2026-09-09_Runpod/results_e5 (no template)/definitions/{}_42.jsonl",
}
OUT_CSV = Path(__file__).resolve().parent / "outputs" / "7_second_pass_rows.csv"


def is_base(name):
    return name.endswith("-pt")


def group_of(dataset, name):
    if dataset == "main":
        return "main base (5)" if is_base(name) else "main instruct (18)"
    return {"E3": "E3 few-shot (5 base)", "E4": "E4 8-word cap (18)", "E5": "E5 no template (18)"}[dataset]


def load(path):
    with open(path, encoding="utf-8") as fh:
        return [json.loads(line) for line in fh if line.strip()]


def edge_counts(job):
    label, path = job
    records = load(path)
    extracted = [dict(r, definition=first_sentence(r.get("definition") or "")) for r in records]
    return label, build_graph(records).number_of_edges(), build_graph(extracted).number_of_edges()


def main():
    stats = collections.defaultdict(collections.Counter)
    parts = collections.defaultdict(collections.Counter)
    base_examples = []
    per_base = {}
    idempotent, colon_end, long_out = True, 0, []
    for dataset, template in PATHS.items():
        for name in NAMES:
            path = ROOT / template.format(name)
            if not path.exists():
                continue
            group = group_of(dataset, name)
            records = [r for r in load(path) if (r.get("definition") or "").strip()]
            removed_here = 0
            for r in records:
                stored = r["definition"]
                out = first_sentence(stored)
                idempotent &= first_sentence(out) == out
                colon_end += out.endswith(":")
                if len(out.split()) >= 60:
                    long_out.append(name)
                c = stats[group]
                c["n"] += 1
                c["differs"] += out != stored.strip()
                removed = cut_removes_content(stored)
                c["removed"] += removed
                removed_here += removed
                c["empty"] += out == ""
                c["w_before"] += len(stored.split())
                c["w_after"] += len(out.split())
                if removed:
                    lines = _cleaned_lines(stored)
                    first = _cut_sentence(lines[0])
                    lead_in = first.endswith(":") and len(lines) > 1
                    parts[group]["lead-in joined"] += lead_in
                    parts[group]["cut inside the first line"] += first != lines[0]
                    parts[group]["later lines dropped"] += len(lines) > (2 if lead_in else 1)
                    if dataset == "main" and is_base(name):
                        base_examples.append((name, stored, out))
            if dataset == "main" and is_base(name):
                per_base[name] = (len(records), removed_here)

    print("1. HOW OFTEN THE RULE CHANGES A RECORD (non-empty stored definitions)")
    print(f"   {'group':24}{'records':>9}{'text differs':>14}{'words removed':>15}{'-> empty':>10}{'words before':>14}{'after':>8}")
    for g, c in stats.items():
        print(f"   {g:24}{c['n']:>9}{100 * c['differs'] / c['n']:>13.1f}%{100 * c['removed'] / c['n']:>14.1f}%"
              f"{100 * c['empty'] / c['n']:>9.1f}%{c['w_before'] / c['n']:>14.1f}{c['w_after'] / c['n']:>8.1f}")
    print("\n   base models, share of records whose words are removed:")
    for name, (n, k) in per_base.items():
        print(f"     {name:16} {k:5} of {n} ({100 * k / n:.1f}%)")

    print("\n2. WHICH PART OF THE RULE CUTS (records with words removed; one record can count in several)")
    for g, c in parts.items():
        print(f"   {g:24} {dict(c)}")
    print(f"\n3. CHECKS: idempotent on every stored text: {idempotent}; extracted texts ending in ':' "
          f"(lead-in without a next line): {colon_end}; extracted texts of 60+ words: {len(long_out)}"
          + (f" {dict(collections.Counter(long_out).most_common(5))}" if long_out else ""))

    print("\n4. GRAPH EDGES BEFORE -> AFTER (same vocabulary, nothing filtered)")
    jobs = [(f"base {n}", str(ROOT / PATHS["main"].format(n))) for n in NAMES if is_base(n)]
    jobs += [(f"E5 {n}", str(ROOT / PATHS["E5"].format(n))) for n in NAMES if (ROOT / PATHS["E5"].format(n)).exists()]
    with Pool(6) as pool:
        for label, before, after in pool.map(edge_counts, jobs):
            print(f"   {label:30} {before:6} -> {after:6}  ({100 * (after / before - 1):+.0f}%)")

    print("\n5. LABELED ROWS WHOSE CONTENT THE RULE CUTS (the second pass)")
    wb = openpyxl.load_workbook(ROOT / "data" / "exp2_labeling_workbook_labeled.xlsx")
    rows = []
    for tab in [s for s in wb.sheetnames if s != "Read me"]:
        ws = wb[tab]
        for r in range(3, 103):
            stored = ws.cell(r, 2).value or ""
            if stored.strip() and cut_removes_content(stored):
                rows.append(dict(tab=tab, row=r, word=ws.cell(r, 1).value, stored=stored.replace("\n", "\\n"),
                                 extracted=first_sentence(stored), first_pass_label=ws.cell(r, 3).value,
                                 first_pass_note=ws.cell(r, 4).value or ""))
    print(f"   rows: {len(rows)}; by tab: {dict(collections.Counter(x['tab'] for x in rows))}")
    print(f"   first-pass label of these rows: {dict(collections.Counter(x['first_pass_label'] for x in rows))}")
    print(f"   rows whose extracted text is empty: {sum(1 for x in rows if x['extracted'] == '')}")
    OUT_CSV.parent.mkdir(exist_ok=True)
    with open(OUT_CSV, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    print(f"   wrote {OUT_CSV.relative_to(ROOT)}")

    print("\n6. TEN RANDOM CUTS IN BASE-MODEL OUTPUTS (stored text shortened)")
    rng = random.Random(7)
    for name, stored, out in rng.sample(base_examples, 10):
        print(f"   [{name}] {stored[:110]!r}\n        -> {out[:110]!r}")


if __name__ == "__main__":
    main()
