"""Create the E2 self-consistency sheet and its key (decision D005).

The researcher re-labels a random sample of their own first-pass rows, blind to the first
labels, at least one day later; 9_selfconsistency_agreement.py then computes agreement.

Sample: 10 random rows from each of the 11 Tier 1/2 models (110 rows, seed fixed below), plus
8 targeted rows that were flagged as borderline in the label review (Qwen3-0.6B imperatives,
Gemma3-270M-pt rows 22 and 31). Targeted rows are analysed separately from the random sample.

Sheet: ID, Word, Model Output, Label (dropdown of the seven labels), Notes. No model names, no
first-pass labels, rows shuffled. Key (first-pass labels and origin of each row) goes to a
separate file that must not be opened before labeling is finished.

Refuses to overwrite existing files, so a filled sheet can never be lost by re-running.
"""
import csv
import random
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data" / "exp2_labeling_workbook_labeled.xlsx"
SHEET = ROOT / "data" / "exp2_selfconsistency_sheet.xlsx"
KEY = ROOT / "data" / "exp2_selfconsistency_KEY_do_not_open_until_done.csv"
SEED = "e2-selfconsistency-2026-10-01"
LABELS = "definition,echo,question,offtopic,restatement,refusal,unclear"
TIER_1_2 = ["Gemma3-27B-pt", "Gemma3-4B-pt", "Gemma3-27B", "Qwen2.5-0.5B",               # Tier 1
            "Gemma3-12B-pt", "Gemma3-12B", "Gemma3-1B-pt", "Gemma3-1B",                  # Tier 2
            "Gemma3-270M-pt", "Gemma3-270M", "Gemma3-4B"]
TARGETED = [("Qwen3-0.6B", 4, "pinpoint"), ("Qwen3-0.6B", 14, "bury"), ("Qwen3-0.6B", 85, "group"),
            ("Qwen3-0.6B", 31, "devise"), ("Qwen3-0.6B", 47, "wait"), ("Qwen3-0.6B", 99, "acknowledge"),
            ("Gemma3-270M-pt", 22, "branch"), ("Gemma3-270M-pt", 31, "devise")]
N_PER_MODEL = 10

for path in (SHEET, KEY):
    if path.exists():
        sys.exit(f"{path.name} already exists; refusing to overwrite (it may hold your labels)")

source = openpyxl.load_workbook(SOURCE)
rng = random.Random(SEED)
picked = []  # (kind, tab, row)
targeted_rows = {(tab, row) for tab, row, _ in TARGETED}
for tab in TIER_1_2:
    candidates = [r for r in range(3, 103) if (tab, r) not in targeted_rows]
    picked += [("random", tab, r) for r in sorted(rng.sample(candidates, N_PER_MODEL))]
for tab, row, word in TARGETED:
    assert source[tab].cell(row, 1).value == word, (tab, row, word, source[tab].cell(row, 1).value)
    picked.append(("targeted", tab, row))
rng.shuffle(picked)

wb = openpyxl.Workbook()
readme = wb.active
readme.title = "Read me"
lines = [
    ("Self-consistency re-label (E2)", True),
    ("", False),
    ("Purpose: measure how consistently you apply your own labels. You label a random sample of your "
     "first-pass rows again, without seeing the first labels.", False),
    ("Start no earlier than 2026-10-02 (at least one day after the first pass ended on 2026-10-01).", False),
    ("", False),
    ("Use the same seven labels (dropdown) and the same rules as the first pass: "
     "research/E2_LABELING_INSTRUCTIONS.md sections 0-3 and 7. Judge the whole output as shown, "
     "exactly as in the first pass. Notes are optional; use the same tags if you add any.", False),
    ("Do not open the key file (data/exp2_selfconsistency_KEY_do_not_open_until_done.csv) and do not "
     "look at the first-pass workbook while you label. Do not reorder or delete rows.", False),
    ("Rows come from several models and are shuffled; the model is not shown on purpose.", False),
    ("When you have labeled every row, save the file and tell the assistant; agreement "
     "(percent and Cohen's kappa) is computed with research/audit_scripts/9_selfconsistency_agreement.py.", False),
]
for i, (text, bold) in enumerate(lines, start=1):
    cell = readme.cell(i, 1, text)
    cell.font = Font(name="Arial", size=11, bold=bold)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
readme.column_dimensions["A"].width = 110

ws = wb.create_sheet("Label these")
header_fill = PatternFill("solid", fgColor="1F2937")
for col, (title, width) in enumerate([("ID", 6), ("Word", 18), ("Model Output", 80), ("Label", 14), ("Notes", 40)], start=1):
    cell = ws.cell(1, col, title)
    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell.fill = header_fill
    ws.column_dimensions[chr(64 + col)].width = width
key_rows = []
for i, (kind, tab, row) in enumerate(picked, start=1):
    word = source[tab].cell(row, 1).value
    output = source[tab].cell(row, 2).value
    first_label = source[tab].cell(row, 3).value
    first_note = source[tab].cell(row, 4).value or ""
    r = i + 1
    ws.cell(r, 1, i)
    ws.cell(r, 2, word)
    ws.cell(r, 3, output)
    for col in (1, 2, 3, 4, 5):
        ws.cell(r, col).font = Font(name="Arial", size=11)
    ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
    key_rows.append(dict(id=i, kind=kind, tab=tab, row=row, word=word,
                         first_pass_label=first_label, first_pass_note=first_note))
ws.freeze_panes = "A2"
validation = DataValidation(type="list", formula1=f'"{LABELS}"', allow_blank=True, showErrorMessage=False)
validation.add(f"D2:D{len(picked) + 1}")
ws.add_data_validation(validation)
wb.save(SHEET)

with open(KEY, "w", newline="", encoding="utf-8-sig") as fh:
    writer = csv.DictWriter(fh, fieldnames=list(key_rows[0].keys()))
    writer.writeheader()
    writer.writerows(key_rows)

print(f"sheet: {SHEET.relative_to(ROOT)}  ({len(picked)} rows: "
      f"{sum(1 for k in key_rows if k['kind'] == 'random')} random + "
      f"{sum(1 for k in key_rows if k['kind'] == 'targeted')} targeted)")
print(f"key  : {KEY.relative_to(ROOT)}  (do not open before labeling is finished)")
per_tab = {}
for k in key_rows:
    if k["kind"] == "random":
        per_tab[k["tab"]] = per_tab.get(k["tab"], 0) + 1
print("random rows per model:", per_tab)
