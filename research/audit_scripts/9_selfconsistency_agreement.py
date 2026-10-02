"""Agreement between the first-pass labels and the blind re-label (decision D005).

Usage (defaults are the real files):
    python research/audit_scripts/9_selfconsistency_agreement.py
    python research/audit_scripts/9_selfconsistency_agreement.py --sheet SHEET.xlsx --key KEY.csv

Reports percent agreement and Cohen's kappa on the random sample (all seven labels, and the
binary definition / not-definition split), by model group and by first-pass class, a confusion
matrix, and every disagreement with both labels. The 8 targeted rows are listed separately.
Notes are not compared.
"""
import argparse
import collections
import csv
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
LABELS = ["definition", "echo", "question", "offtopic", "restatement", "refusal", "unclear"]
sys.stdout.reconfigure(encoding="utf-8")


def kappa(pairs):
    n = len(pairs)
    if n == 0:
        return float("nan")
    po = sum(a == b for a, b in pairs) / n
    first, second = collections.Counter(a for a, _ in pairs), collections.Counter(b for _, b in pairs)
    pe = sum(first[k] * second[k] for k in set(first) | set(second)) / (n * n)
    return float("nan") if pe == 1 else (po - pe) / (1 - pe)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--sheet", default=str(ROOT / "data" / "exp2_selfconsistency_sheet.xlsx"))
    parser.add_argument("--key", default=str(ROOT / "data" / "exp2_selfconsistency_KEY_do_not_open_until_done.csv"))
    args = parser.parse_args()

    key = {int(r["id"]): r for r in csv.DictReader(open(args.key, encoding="utf-8-sig"))}
    ws = openpyxl.load_workbook(args.sheet)["Label these"]
    second = {}
    for r in range(2, ws.max_row + 1):
        rid, label = ws.cell(r, 1).value, ws.cell(r, 4).value
        if rid is not None:
            second[int(rid)] = (str(label).strip().lower() if label not in (None, "") else None,
                                ws.cell(r, 5).value or "", ws.cell(r, 3).value or "")
    missing = [i for i in key if second.get(i, (None,))[0] is None]
    invalid = [(i, second[i][0]) for i in key if second.get(i, (None,))[0] not in (None, *LABELS)]
    print(f"rows {len(key)}; unlabeled {len(missing)} {missing[:10]}; invalid labels {invalid}")
    if missing or invalid:
        sys.exit("fix the sheet first")

    def pairs_of(kind):
        return [(key[i]["first_pass_label"], second[i][0], i) for i in key if key[i]["kind"] == kind]

    rnd = pairs_of("random")
    pp = [(a, b) for a, b, _ in rnd]
    binary = [("definition" if a == "definition" else "other", "definition" if b == "definition" else "other") for a, b in pp]
    agree = sum(a == b for a, b in pp)
    print(f"\nRANDOM SAMPLE (n={len(rnd)})")
    print(f"  all seven labels: agreement {agree}/{len(rnd)} = {100 * agree / len(rnd):.1f}%, Cohen's kappa = {kappa(pp):.2f}")
    bagree = sum(a == b for a, b in binary)
    print(f"  definition vs not: agreement {bagree}/{len(rnd)} = {100 * bagree / len(rnd):.1f}%, Cohen's kappa = {kappa(binary):.2f}")
    groups = {"base models (5)": lambda t: t.endswith("-pt"), "Gemma3-270M": lambda t: t == "Gemma3-270M",
              "other instruct (5)": lambda t: not t.endswith("-pt") and t != "Gemma3-270M"}
    for g, test in groups.items():
        sub = [(a, b) for a, b, i in rnd if test(key[i]["tab"])]
        if sub:
            print(f"  {g:20} n={len(sub):3}  agreement {100 * sum(a == b for a, b in sub) / len(sub):5.1f}%  kappa {kappa(sub):.2f}")

    print("\n  by first-pass class (how often the same label came back):")
    for lab in LABELS:
        sub = [(a, b) for a, b in pp if a == lab]
        if sub:
            print(f"    {lab:12} first-pass n={len(sub):3}  same again {sum(a == b for a, b in sub):3} ({100 * sum(a == b for a, b in sub) / len(sub):5.1f}%)")
    print("\n  confusion (rows = first pass, columns = second pass)")
    print("    " + " " * 12 + "".join(f"{l[:6]:>8}" for l in LABELS))
    cm = collections.Counter(pp)
    for a in LABELS:
        if any(cm[(a, b)] for b in LABELS):
            print(f"    {a:12}" + "".join(f"{cm[(a, b)]:>8}" for b in LABELS))

    print("\n  disagreements (random sample):")
    for a, b, i in sorted(rnd, key=lambda x: x[2]):
        if a != b:
            k = key[i]
            print(f"    id {i:3} [{k['tab']} row {k['row']}, {k['word']}] first={a} ({k['first_pass_note'] or '-'}) second={b} ({second[i][1] or '-'})"
                  f"\n        {str(second[i][2]).replace(chr(10), ' ')[:140]!r}")
    print("\nTARGETED ROWS (borderline in the first pass)")
    for a, b, i in sorted(pairs_of("targeted"), key=lambda x: x[2]):
        k = key[i]
        print(f"    id {i:3} [{k['tab']} row {k['row']}, {k['word']}] first={a:11} second={b:11} {'same' if a == b else 'CHANGED'}  "
              f"{str(second[i][2]).replace(chr(10), ' ')[:70]!r}")


if __name__ == "__main__":
    main()
