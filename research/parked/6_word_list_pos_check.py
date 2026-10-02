"""PARKED 2026-10-01: out of scope by the researcher's decision (DECISION_LOG D017); no document relies on it.

How often does the word list ask for a part of speech that is rare for that word?

Audit script (2026-10-01). The E2 labeling found 72 rows where the model defined a different
part of speech than the prompt asked for ("wrong-pos"), concentrated on 26 of the 100 sampled
words. This script measures, for every entry of the 3,000-word list, how much of the lemma's
usage in WordNet's SemCor counts falls on the requested part of speech, and checks that the
measure separates the 26 flagged words from the other 74.

share = SemCor count of the lemma under the requested POS / its count over all POS.
Entries with no SemCor count at all have share = None (nothing to judge).
"""
import collections
import json
import sys
from pathlib import Path

import openpyxl
from nltk.corpus import wordnet as wn

ROOT = Path(__file__).resolve().parents[2]  # research/parked/ -> repo root
sys.stdout.reconfigure(encoding="utf-8")
words = json.load(open(ROOT / "data" / "sample_words" / "word_list_3k_v1.json", encoding="utf-8"))["words"]


def counts(lemma: str) -> dict:
    """SemCor count of `lemma` under each POS (adjective satellites counted as adjectives)."""
    out = collections.Counter()
    key = lemma.replace(" ", "_").lower()
    for syn in wn.synsets(key):
        pos = "a" if syn.pos() == "s" else syn.pos()
        for lem in syn.lemmas():
            if lem.name().lower() == key:
                out[pos] += lem.count()
    return out


def share(entry: dict):
    c = counts(entry["lemma"])
    total = sum(c.values())
    return (c[entry["pos"]] / total if total else None), total


shares = [(e, *share(e)) for e in words]
known = [(e, s, t) for e, s, t in shares if s is not None]
print(f"entries {len(words)}; with any SemCor count {len(known)}; none {len(words) - len(known)}")
for thr in (0.05, 0.10, 0.20):
    low = [x for x in known if x[1] < thr]
    print(f"  requested POS has share < {thr:.2f}: {len(low)} entries "
          f"({100 * len(low) / len(known):.1f}% of entries with counts; "
          f"by POS {dict(collections.Counter(x[0]['pos'] for x in low))})")
zero = [x for x in known if x[1] == 0]
print(f"  share exactly 0 (lemma is used only under other POS in SemCor): {len(zero)} "
      f"({100 * len(zero) / len(known):.1f}%)")

# do the 26 words the labeler flagged as wrong-pos stand out?
wb = openpyxl.load_workbook(ROOT / "data" / "exp2_labeling_workbook_labeled.xlsx")
flagged = set()
sample = {}
for tab in [n for n in wb.sheetnames if n != "Read me"]:
    ws = wb[tab]
    for r in range(3, 103):
        w, pos = ws.cell(r, 1).value, ws.cell(r, 5).value
        sample[(w, pos)] = True
        if "wrong-pos" in str(ws.cell(r, 4).value or "").lower():
            flagged.add((w, pos))
code = {"noun": "n", "verb": "v", "adj": "a"}
lookup = {(e["lemma"], e["pos"]): (s, t) for e, s, t in shares}
fl, ok = [], []
for (w, pos) in sample:
    key = (w, code.get(pos.split("/")[0]))
    if key not in lookup:
        continue
    (fl if (w, pos) in flagged else ok).append((w, pos, *lookup[key]))
f_s = [s for _, _, s, _ in fl if s is not None]
o_s = [s for _, _, s, _ in ok if s is not None]
print(f"\nthe sampled 100 words: flagged wrong-pos {len(fl)} entries, not flagged {len(ok)}")
print(f"  mean share   flagged {sum(f_s) / len(f_s):.2f} (n={len(f_s)})   not flagged {sum(o_s) / len(o_s):.2f} (n={len(o_s)})")
for thr in (0.05, 0.10, 0.20):
    a = sum(1 for s in f_s if s < thr); b = sum(1 for s in o_s if s < thr)
    print(f"  share < {thr:.2f}: flagged {a}/{len(f_s)}, not flagged {b}/{len(o_s)}")
print("\nlowest shares among flagged entries (word, POS asked, share, SemCor total):")
for w, pos, s, t in sorted(fl, key=lambda x: (x[2] is None, x[2] if x[2] is not None else 9))[:12]:
    print(f"  {w:14} {pos:6} share={'n/a' if s is None else f'{s:.2f}'} total={t}")
print("\nunflagged entries with share < 0.10 (the labeler did not note a POS problem for these):")
for w, pos, s, t in sorted(ok, key=lambda x: x[2] if x[2] is not None else 9)[:10]:
    if s is not None and s < 0.10:
        print(f"  {w:14} {pos:6} share={s:.2f} total={t}")
